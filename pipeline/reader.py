"""
文档读取：支持 .docx / .doc / .txt / .pdf
"""

import re
from pathlib import Path


def read_document(filepath: str) -> str:
    """读取文档内容，返回纯文本。

    支持格式: .docx / .doc / .txt / .pdf
    失败时抛出异常。
    """
    path = Path(filepath)
    ext = path.suffix.lower()

    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")

        doc = Document(filepath)
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    tables_text.append(" | ".join(cells))
        all_text = "\n".join(paras)
        if tables_text:
            all_text += "\n\n【表格内容】\n" + "\n".join(tables_text)
        return all_text

    elif ext == ".txt":
        with open(filepath, "r", encoding="utf-8") as f:
            return f.read()

    elif ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("请安装 pypdf: pip install pypdf")

        reader = PdfReader(filepath)
        texts = [page.extract_text() for page in reader.pages if page.extract_text()]
        return "\n".join(texts)

    elif ext == ".doc":
        try:
            import olefile
        except ImportError:
            raise ImportError("请安装 olefile: pip install olefile")

        try:
            ole = olefile.OleFileIO(filepath)
            wd = ole.openstream("WordDocument").read()
            ole.close()
        except Exception as e:
            raise ValueError(f"读取 .doc 文件失败（olefile）: {e}")

        # Word Binary 格式存储 UTF-16LE 文本
        text = wd.decode("utf-16le", errors="replace")
        lines = []
        for line in text.split("\n"):
            cleaned = line.strip()
            chinese_count = sum(1 for c in cleaned if "一" <= c <= "鿿")
            if chinese_count >= 5 and len(cleaned) > 5:
                lines.append(cleaned)
        result = "\n".join(lines)
        result = re.sub(r"[^一-鿿　-〿＀-￯0-9a-zA-Z\s，。、；：！？（）【】\"\"''《》—…·\n—-]", "", result)
        result = re.sub(r"\n{3,}", "\n\n", result).strip()
        return result

    else:
        raise ValueError(f"不支持的文件格式: {ext}，支持 .docx / .doc / .txt / .pdf")
